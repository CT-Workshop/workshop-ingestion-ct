from __future__ import annotations

import hashlib
import io
import logging
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

logger = logging.getLogger(__name__)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_metadata(
    *,
    content: bytes,
    original_name: str,
    declared_content_type: str | None,
) -> dict[str, Any]:
    """
    Best-effort metadata extraction for supported office formats.
    """
    suffix = Path(original_name).suffix.lower()
    base: dict[str, Any] = {
        "original_filename": original_name,
        "declared_content_type": declared_content_type,
        "extension": suffix,
        "sha256": sha256_bytes(content),
        "size_bytes": len(content),
    }

    if suffix == ".pdf":
        base.update(_pdf_meta(content))
    elif suffix == ".docx":
        base.update(_docx_meta(content))
    elif suffix == ".xlsx":
        base.update(_xlsx_meta(content))
    else:
        base["parser"] = "none"

    return base


def _pdf_meta(content: bytes) -> dict[str, Any]:
    out: dict[str, Any] = {"parser": "pypdf"}
    try:
        reader = PdfReader(io.BytesIO(content))
        out["page_count"] = len(reader.pages)
        meta = reader.metadata
        if meta:
            out["pdf_title"] = meta.get("/Title")
            out["pdf_author"] = meta.get("/Author")
    except Exception as exc:  # noqa: BLE001 — best-effort extraction
        logger.warning("pdf metadata extraction failed: %s", exc)
        out["pdf_error"] = str(exc)
    return out


def _docx_meta(content: bytes) -> dict[str, Any]:
    out: dict[str, Any] = {"parser": "python-docx"}
    try:
        doc = DocxDocument(io.BytesIO(content))
        core = doc.core_properties
        out["title"] = core.title
        out["author"] = core.author
        out["subject"] = core.subject
        out["paragraph_count"] = len(doc.paragraphs)
    except Exception as exc:  # noqa: BLE001
        logger.warning("docx metadata extraction failed: %s", exc)
        out["docx_error"] = str(exc)
    return out


def _xlsx_meta(content: bytes) -> dict[str, Any]:
    out: dict[str, Any] = {"parser": "openpyxl"}
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        out["sheet_names"] = list(wb.sheetnames)
        out["sheet_count"] = len(wb.sheetnames)
        wb.close()
    except Exception as exc:  # noqa: BLE001
        logger.warning("xlsx metadata extraction failed: %s", exc)
        out["xlsx_error"] = str(exc)
    return out
